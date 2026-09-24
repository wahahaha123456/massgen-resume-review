"""解析器单元验证：造 txt/docx/pdf 三种文件 + 错误用例，不调 LLM。"""

import io
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, r"d:\MassGen")
from api.resume_parser import ResumeParseError, parse_resume_file

RESUME = (
    "张三，2024年某大学计算机专业毕业。技能：Python、Django、MySQL、Redis、Docker。"
    "项目经历：校园二手交易平台，负责后端开发，日均活跃500人。"
)

results = []

# 1. txt
txt_text = parse_resume_file("resume.txt", RESUME.encode("utf-8"))
results.append(("txt", len(txt_text) > 10, txt_text[:40]))

# 2. docx
import docx
doc = docx.Document()
doc.add_heading("张三 - 简历", level=1)
doc.add_paragraph("2024年某大学计算机专业毕业。技能：Python、Django、MySQL、Redis、Docker。")
doc.add_paragraph("项目经历：校园二手交易平台，负责后端开发，日均活跃500人。")
# 加一个表格，验证表格提取
table = doc.add_table(rows=2, cols=2)
table.rows[0].cells[0].text = "技能"
table.rows[0].cells[1].text = "熟练度"
table.rows[1].cells[0].text = "Python"
table.rows[1].cells[1].text = "熟练"
buf = io.BytesIO()
doc.save(buf)
docx_text = parse_resume_file("resume.docx", buf.getvalue())
results.append(("docx", "张三" in docx_text and "Python" in docx_text and "熟练" in docx_text, docx_text[:60].replace("\n", " | ")))

# 3. pdf (reportlab)
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import os

# reportlab 默认字体不支持中文，找系统中文字体注册
font_name = "Helvetica"
candidates = [
    r"C:\Windows\Fonts\msyh.ttc",
    r"C:\Windows\Fonts\simsun.ttc",
    r"C:\Windows\Fonts\simhei.ttf",
]
for fp in candidates:
    if os.path.exists(fp):
        try:
            pdfmetrics.registerFont(TTFont("CJK", fp))
            font_name = "CJK"
            break
        except Exception:
            continue

buf = io.BytesIO()
c = canvas.Canvas(buf, pagesize=A4)
c.setFont(font_name, 12)
y = 800
for line in ["Zhang San Resume", "2024 Computer Science Graduate", "Skills: Python, Django, MySQL, Redis, Docker",
             "Project: campus marketplace backend, 500 DAU"]:
    c.drawString(72, y, line)
    y -= 24
c.save()
pdf_text = parse_resume_file("resume.pdf", buf.getvalue())
results.append(("pdf", "Python" in pdf_text and "500 DAU" in pdf_text, pdf_text[:80].replace("\n", " | ")))

# 4. 错误用例：不支持的扩展名
try:
    parse_resume_file("resume.xlsx", b"fake")
    results.append(("bad-ext", False, "未拒绝 xlsx"))
except ResumeParseError as e:
    results.append(("bad-ext", "不支持" in str(e), str(e)))

# 5. 错误用例：空文件
try:
    parse_resume_file("resume.txt", b"")
    results.append(("empty", False, "未拒绝空文件"))
except ResumeParseError as e:
    results.append(("empty", "为空" in str(e), str(e)))

# 6. 错误用例：损坏 PDF
try:
    parse_resume_file("resume.pdf", b"not a real pdf content here")
    results.append(("bad-pdf", False, "未拒绝损坏PDF"))
except ResumeParseError as e:
    results.append(("bad-pdf", True, str(e)[:60]))

print("=" * 70)
all_pass = True
for name, ok, detail in results:
    mark = "PASS" if ok else "FAIL"
    if not ok:
        all_pass = False
    print(f"[{mark}] {name:10s} {detail}")
print("=" * 70)
print("ALL PASS" if all_pass else "SOME FAILED")
sys.exit(0 if all_pass else 1)
