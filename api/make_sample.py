"""造一个中文 docx 测试简历，供上传端到端测试使用。"""

import sys

sys.path.insert(0, r"d:\MassGen")

import docx

doc = docx.Document()
doc.add_heading("张三 - 个人简历", level=1)

doc.add_heading("教育背景", level=2)
doc.add_paragraph("2020-2024 某大学 计算机科学与技术 本科")

doc.add_heading("专业技能", level=2)
doc.add_paragraph("编程语言：Python（熟练）、Java（了解）")
doc.add_paragraph("框架与中间件：Django、MySQL、Redis、Docker、Nginx")
doc.add_paragraph("工具：Git、Linux、Postman")

doc.add_heading("项目经历", level=2)
p = doc.add_paragraph()
p.add_run("1. 校园二手交易平台（2023.03-2023.12） 后端开发").bold = True
doc.add_paragraph("用 Django 实现用户注册登录、商品发布与检索模块，使用 MySQL 存储、Redis 缓存热点商品。")
doc.add_paragraph("平台上线后日均活跃用户约 500 人，商品列表接口响应时间从 500ms 优化到 120ms。")

p = doc.add_paragraph()
p.add_run("2. 某科技公司实习（2024.03-2024.06） 数据开发实习生").bold = True
doc.add_paragraph("用 Python 编写数据清洗脚本，处理业务数据约 10 万条，输出标准化报表。")

doc.add_heading("自我评价", level=2)
doc.add_paragraph("热爱后端开发，有良好的编码习惯和团队协作能力。")

out = r"d:\MassGen\api\sample_resume.docx"
doc.save(out)
print(f"saved -> {out}")
