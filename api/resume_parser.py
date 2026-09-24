"""简历文件解析：按扩展名分发，第三方库延迟导入。

支持格式：.pdf / .docx / .txt / .md
PDF 表格类复杂布局可能解析乱序，当前为简单版（逐页提取文本）。
"""

from __future__ import annotations

import io
from pathlib import Path

# 前后端一致的扩展名白名单（前端 accept 也用它）
SUPPORTED_EXTS: frozenset[str] = frozenset({".pdf", ".docx", ".txt", ".md"})

MAX_FILE_BYTES = 10 * 1024 * 1024  # 10MB，防止超大文件打爆子进程 prompt


class ResumeParseError(ValueError):
    """简历文件解析失败（格式不支持/损坏/无文本）。"""


def get_ext(filename: str) -> str:
    return Path(filename).suffix.lower()


def extract_pdf(data: bytes) -> str:
    """PyPDF2 逐页提取文本。"""
    from PyPDF2 import PdfReader  # 延迟导入

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # 损坏/加密 PDF 会抛各类异常
        raise ResumeParseError(f"PDF 无法读取（可能已损坏或加密）：{exc}") from exc

    pages: list[str] = []
    for i, page in enumerate(reader.pages):
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:
            raise ResumeParseError(f"PDF 第 {i + 1} 页解析失败：{exc}") from exc

    return "\n\n".join(p.strip() for p in pages if p.strip())


def extract_docx(data: bytes) -> str:
    """python-docx 提取段落文本（含表格单元格）。"""
    import docx  # 延迟导入

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeParseError(f"DOCX 无法读取（可能已损坏，注意不支持旧版 .doc）：{exc}") from exc

    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]

    # 表格内容也抽出来，避免表格里的经历/技能丢失
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                # 去重连续相同单元格（合并单元格会重复出现）
                deduped: list[str] = []
                for cell in cells:
                    if not deduped or cell != deduped[-1]:
                        deduped.append(cell)
                parts.append(" | ".join(deduped))

    return "\n".join(parts)


def extract_text(data: bytes) -> str:
    """txt/md 直接按 UTF-8 解码（失败回退 GBK，兼容 Windows 导出的旧文件）。"""
    for encoding in ("utf-8", "utf-8-sig", "gbk"):
        try:
            return data.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    raise ResumeParseError("文本文件编码无法识别（请另存为 UTF-8）")


def parse_resume_file(filename: str, data: bytes) -> str:
    """统一入口：校验扩展名/大小 → 分发解析 → 返回纯文本。"""
    ext = get_ext(filename)
    if ext not in SUPPORTED_EXTS:
        allowed = "、".join(sorted(SUPPORTED_EXTS))
        raise ResumeParseError(f"不支持的文件类型 {ext or '（无扩展名）'}，仅支持：{allowed}")

    if len(data) == 0:
        raise ResumeParseError("文件为空")
    if len(data) > MAX_FILE_BYTES:
        raise ResumeParseError(f"文件过大（{len(data) / 1024 / 1024:.1f}MB），上限 10MB")

    if ext == ".pdf":
        text = extract_pdf(data)
    elif ext == ".docx":
        text = extract_docx(data)
    else:
        text = extract_text(data)

    text = text.strip()
    if not text:
        raise ResumeParseError("未能从文件中提取到文本（PDF 可能是扫描件/纯图片）")
    return text
